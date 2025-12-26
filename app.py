import os
import json
import io
from datetime import datetime, timedelta
from flask import Flask, render_template, request, redirect, url_for, flash, jsonify, send_file
from flask_sqlalchemy import SQLAlchemy
from flask_login import LoginManager, UserMixin, login_user, login_required, logout_user, current_user
from werkzeug.security import generate_password_hash, check_password_hash
from werkzeug.utils import secure_filename
from functools import wraps
import openpyxl 

app = Flask(__name__)
app.secret_key = 'bi_mat_cua_teo_v7_final' 

# --- KẾT NỐI DATABASE ---
db_url = os.environ.get('DATABASE_URL')
if db_url and db_url.startswith("postgres://"):
    db_url = db_url.replace("postgres://", "postgresql://", 1)
app.config['SQLALCHEMY_DATABASE_URI'] = db_url
app.config['SQLALCHEMY_TRACK_MODIFICATIONS'] = False

db = SQLAlchemy(app)
login_manager = LoginManager()
login_manager.init_app(app)
login_manager.login_view = 'login'

# --- MODELS ---
class User(UserMixin, db.Model):
    id = db.Column(db.Integer, primary_key=True)
    username = db.Column(db.String(100), unique=True, nullable=False)
    password = db.Column(db.String(200), nullable=False)
    role = db.Column(db.String(20), default='user')

class AppTable(db.Model):
    id = db.Column(db.Integer, primary_key=True)
    name = db.Column(db.String(100), nullable=False, unique=True)

class TableColumn(db.Model):
    id = db.Column(db.Integer, primary_key=True)
    name = db.Column(db.String(100), nullable=False)
    order_index = db.Column(db.Integer, default=0)
    user_id = db.Column(db.Integer, db.ForeignKey('user.id'), nullable=False)
    table_id = db.Column(db.Integer, db.ForeignKey('app_table.id'), nullable=False)

class DataRow(db.Model):
    id = db.Column(db.Integer, primary_key=True)
    content = db.Column(db.JSON, nullable=False) 
    created_by = db.Column(db.Integer, db.ForeignKey('user.id'), nullable=False)
    table_id = db.Column(db.Integer, db.ForeignKey('app_table.id'), nullable=False)
    created_at = db.Column(db.DateTime, default=datetime.utcnow) 

@login_manager.user_loader
def load_user(user_id):
    return User.query.get(int(user_id))

def admin_required(f):
    @wraps(f)
    def decorated_function(*args, **kwargs):
        if not current_user.is_authenticated or current_user.role != 'admin':
            flash('Chỉ dành cho Admin!', 'error')
            return redirect(url_for('index'))
        return f(*args, **kwargs)
    return decorated_function

# --- ROUTES CHÍNH ---

@app.route('/')
@app.route('/table/<int:table_id>')
@login_required
def index(table_id=None):
    all_tables = AppTable.query.order_by(AppTable.id).all()
    
    # Tự tạo bảng mặc định nếu chưa có
    if not all_tables:
        default_table = AppTable(name="DATA XƯỞNG THÉP")
        db.session.add(default_table)
        db.session.commit()
        return redirect(url_for('index'))

    current_table = None
    if table_id:
        current_table = AppTable.query.get(table_id)
    if not current_table:
        current_table = all_tables[0]
        return redirect(url_for('index', table_id=current_table.id))

    # Lấy cột
    columns = TableColumn.query.filter_by(user_id=current_user.id, table_id=current_table.id).order_by(TableColumn.order_index).all()
    
    # Tạo cột mẫu nếu bảng trống
    if not columns:
        defaults = ["Sản phẩm", "Mác thép", "Bộ gap", "Nhiệt lò nung", "Cơ tính"]
        for idx, name in enumerate(defaults):
            db.session.add(TableColumn(name=name, order_index=idx, user_id=current_user.id, table_id=current_table.id))
        db.session.commit()
        return redirect(url_for('index', table_id=current_table.id))

    # Lấy dữ liệu (Load hết, không phân trang)
    rows = DataRow.query.filter_by(created_by=current_user.id, table_id=current_table.id).order_by(DataRow.id.desc()).all()
    data_map = {row.id: row.content for row in rows}

    # Lấy user cho Admin Panel
    all_users = []
    if current_user.role == 'admin':
        all_users = User.query.order_by(User.role, User.id).all()

    return render_template('dashboard.html', 
                           columns=columns, 
                           rows=rows, 
                           data_map=data_map, 
                           all_users=all_users,
                           all_tables=all_tables,
                           current_table=current_table)

# --- XUẤT EXCEL ---
@app.route('/export_excel/<int:table_id>')
@login_required
def export_excel(table_id):
    table = AppTable.query.get_or_404(table_id)
    columns = TableColumn.query.filter_by(user_id=current_user.id, table_id=table.id).order_by(TableColumn.order_index).all()
    rows = DataRow.query.filter_by(created_by=current_user.id, table_id=table.id).order_by(DataRow.id.desc()).all()

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Du Lieu"
    
    # Header
    headers = ["ID", "Ngày tạo"] + [col.name for col in columns]
    ws.append(headers)

    # Data
    for row in rows:
        vn_time = row.created_at + timedelta(hours=7) if row.created_at else datetime.now()
        row_data = [row.id, vn_time.strftime('%d/%m/%Y %H:%M')]
        for col in columns:
            cell_value = row.content.get(col.name, '')
            if isinstance(cell_value, list):
                cell_value = "\n".join(cell_value)
            row_data.append(cell_value)
        ws.append(row_data)

    output = io.BytesIO()
    wb.save(output)
    output.seek(0)
    
    filename = f"{table.name}_{datetime.now().strftime('%Y%m%d')}.xlsx"
    return send_file(output, download_name=filename, as_attachment=True, mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')

# --- TÍNH NĂNG BACKUP JSON (TOÀN BỘ CÁC BẢNG) ---
@app.route('/backup_json')
@login_required
def backup_json():
    # 1. Tạo vỏ bọc cho file backup
    backup_data = {
        "version": "2.0", # Nâng version lên chơi
        "timestamp": datetime.utcnow().strftime('%Y-%m-%d %H:%M:%S'),
        "user": current_user.username,
        "tables": [] # Chỗ này sẽ chứa danh sách các bảng
    }

    # 2. Lấy danh sách TẤT CẢ các mục (Menu) đang có trong hệ thống
    all_system_tables = AppTable.query.all()
    
    # 3. Duyệt qua từng bảng để gom dữ liệu của User
    for table in all_system_tables:
        # A. Lấy Cột của user trong bảng này
        cols = TableColumn.query.filter_by(
            user_id=current_user.id, 
            table_id=table.id
        ).order_by(TableColumn.order_index).all()
        
        # B. Lấy Dòng dữ liệu của user trong bảng này
        rows = DataRow.query.filter_by(
            created_by=current_user.id, 
            table_id=table.id
        ).order_by(DataRow.id).all()
        
        # C. Chỉ backup nếu bảng này user CÓ DỮ LIỆU (Có cột hoặc có dòng)
        # (Để tránh file backup bị rác bởi mấy bảng trống trơn)
        if cols or rows:
            table_data = {
                "name": table.name, # Tên bảng (VD: Xưởng Thép, Kho A...)
                "columns": [{"name": c.name, "order": c.order_index} for c in cols],
                "rows": []
            }
            
            for r in rows:
                table_data["rows"].append({
                    "content": r.content,
                    "created_at": r.created_at.strftime('%Y-%m-%d %H:%M:%S') if r.created_at else None
                })
            
            # Đẩy dữ liệu bảng này vào danh sách tổng
            backup_data["tables"].append(table_data)

    # 4. Xuất file JSON
    json_str = json.dumps(backup_data, indent=4, ensure_ascii=False)
    output = io.BytesIO()
    output.write(json_str.encode('utf-8'))
    output.seek(0)
    
    # Đặt tên file: backup_tenuser_ngay_gio.json
    filename = f"backup_{current_user.username}_FULL_{datetime.now().strftime('%Y%m%d_%H%M')}.json"
    
    return send_file(
        output, 
        download_name=filename, 
        as_attachment=True, 
        mimetype='application/json'
    )
# --- TÍNH NĂNG RESTORE JSON (ĐÃ FIX LỖI SO SÁNH JSON) ---
@app.route('/restore_json', methods=['POST'])
@login_required
def restore_json():
    if 'file' not in request.files:
        flash('Chưa chọn file anh hai ơi!', 'error')
        return redirect(url_for('index'))
    
    file = request.files['file']
    if file.filename == '':
        flash('Tên file trống trơn à!', 'error')
        return redirect(url_for('index'))

    if file:
        try:
            data = json.load(file)
            added_count = 0
            skipped_count = 0
            
            for tbl_data in data.get('tables', []):
                # 1. Tìm bảng trong hệ thống
                sys_table = AppTable.query.filter_by(name=tbl_data['name']).first()
                if not sys_table: continue 

                # 2. Khôi phục Cột (Giữ nguyên)
                for col_data in tbl_data['columns']:
                    exists = TableColumn.query.filter_by(user_id=current_user.id, table_id=sys_table.id, name=col_data['name']).first()
                    if not exists:
                        db.session.add(TableColumn(name=col_data['name'], order_index=col_data['order'], user_id=current_user.id, table_id=sys_table.id))

                # --- ĐOẠN NÀY LÀ KHÁC BIỆT (CHIẾN THUẬT MỚI) ---
                
                # A. Lấy toàn bộ dữ liệu hiện có của user trong bảng này ra trước
                # (Để Python so sánh trong bộ nhớ, tránh bắt Database làm việc khó)
                current_rows = DataRow.query.filter_by(table_id=sys_table.id, created_by=current_user.id).all()
                
                # B. Tạo một danh sách chứa các nội dung (content) đang có
                existing_contents = [row.content for row in current_rows]

                # C. Duyệt qua file backup
                for row_data in tbl_data['rows']:
                    content_to_check = row_data['content']
                    
                    # D. So sánh bằng Python (Python so sánh 2 cái dict cực chuẩn)
                    if content_to_check in existing_contents:
                        skipped_count += 1
                        continue # Bỏ qua, đi tới dòng tiếp theo

                    # E. Nếu chưa có thì thêm mới
                    c_at = datetime.utcnow()
                    if row_data.get('created_at'):
                        try: c_at = datetime.strptime(row_data['created_at'], '%Y-%m-%d %H:%M:%S')
                        except: pass
                    
                    new_row = DataRow(
                        content=content_to_check, 
                        created_by=current_user.id, 
                        table_id=sys_table.id, 
                        created_at=c_at
                    )
                    db.session.add(new_row)
                    
                    # Thêm vào danh sách tạm để check cho các dòng sau trong cùng file backup
                    existing_contents.append(content_to_check)
                    added_count += 1
            
            db.session.commit()
            flash(f'Xong! Đã thêm {added_count} dòng. Bỏ qua {skipped_count} dòng trùng lặp.', 'success')
            
        except Exception as e:
            # In lỗi ra terminal để dễ debug nếu có
            print(f"Lỗi restore: {e}")
            flash(f'Lỗi khi đọc file: {str(e)}', 'error')

    return redirect(url_for('index'))

# --- QUẢN LÝ BẢNG (ADMIN) ---
@app.route('/admin/add_table', methods=['POST'])
@login_required
@admin_required
def add_table():
    name = request.form.get('table_name')
    if name:
        exists = AppTable.query.filter_by(name=name).first()
        if not exists:
            db.session.add(AppTable(name=name))
            db.session.commit()
            flash(f'Đã thêm mục: {name}', 'success')
        else: flash('Tên đã tồn tại!', 'error')
    return redirect(url_for('index'))

@app.route('/admin/edit_table', methods=['POST'])
@login_required
@admin_required
def edit_table():
    table_id = request.form.get('table_id')
    new_name = request.form.get('new_name')
    if table_id and new_name:
        table = AppTable.query.get(table_id)
        if table:
            table.name = new_name
            db.session.commit()
            flash('Đổi tên thành công!', 'success')
    return redirect(url_for('index', table_id=table_id))

@app.route('/admin/delete_table/<int:table_id>')
@login_required
@admin_required
def delete_table(table_id):
    table = AppTable.query.get(table_id)
    if table:
        DataRow.query.filter_by(table_id=table.id).delete()
        TableColumn.query.filter_by(table_id=table.id).delete()
        db.session.delete(table)
        db.session.commit()
        flash('Đã xóa bảng!', 'success')
    return redirect(url_for('index'))

# --- CÁC ROUTE DATA ---
@app.route('/add_column', methods=['POST'])
@login_required
def add_column():
    table_id = request.form.get('table_id')
    col_name = request.form.get('col_name')
    if col_name and table_id:
        exists = TableColumn.query.filter_by(name=col_name, user_id=current_user.id, table_id=table_id).first()
        if not exists:
            max = db.session.query(db.func.max(TableColumn.order_index)).filter_by(user_id=current_user.id, table_id=table_id).scalar() or 0
            db.session.add(TableColumn(name=col_name, order_index=max+1, user_id=current_user.id, table_id=table_id))
            db.session.commit()
            flash(f'Đã thêm cột: {col_name}', 'success')
        else: flash('Trùng tên cột!', 'error')
    return redirect(url_for('index', table_id=table_id))

@app.route('/delete_column/<int:id>')
@login_required
def delete_column(id):
    col = TableColumn.query.get(id)
    if col and col.user_id == current_user.id:
        tid = col.table_id
        db.session.delete(col)
        db.session.commit()
        return redirect(url_for('index', table_id=tid))
    return redirect(url_for('index'))

@app.route('/save_row', methods=['POST'])
@login_required
def save_row():
    row_id = request.form.get('row_id')
    table_id = request.form.get('table_id')
    columns = TableColumn.query.filter_by(user_id=current_user.id, table_id=table_id).all()
    row_data = {}
    for col in columns:
        raw_val = request.form.get(f'field_{col.id}', '')
        if '\n' in raw_val:
            lines = [line.strip() for line in raw_val.split('\n') if line.strip()]
            row_data[col.name] = lines
        else: row_data[col.name] = raw_val.strip()
    
    if row_id:
        row = DataRow.query.get(row_id)
        if row and row.created_by == current_user.id:
            row.content = row_data
            db.session.commit()
            flash('Đã cập nhật!', 'success')
    else:
        db.session.add(DataRow(content=row_data, created_by=current_user.id, table_id=table_id))
        db.session.commit()
        flash('Đã thêm dòng!', 'success')
    return redirect(url_for('index', table_id=table_id))

@app.route('/batch_update_columns', methods=['POST'])
@login_required
def batch_update_columns():
    try:
        data = request.json.get('columns', [])
        if not data: return jsonify({'status': 'success'})
        first_col = TableColumn.query.get(data[0]['id'])
        current_table_id = first_col.table_id
        my_rows = DataRow.query.filter_by(created_by=current_user.id, table_id=current_table_id).all()
        for item in data:
            col = TableColumn.query.get(item.get('id'))
            if col and col.user_id == current_user.id:
                new_name = item.get('name')
                if col.name != new_name:
                    existing = TableColumn.query.filter_by(name=new_name, user_id=current_user.id, table_id=current_table_id).filter(TableColumn.id != col.id).first()
                    if existing: return jsonify({'status': 'error', 'message': f'Tên {new_name} bị trùng!'})
                    old_name = col.name
                    col.name = new_name
                    for row in my_rows:
                        if row.content and old_name in row.content:
                            updated = dict(row.content)
                            updated[new_name] = updated.pop(old_name)
                            row.content = updated
                col.order_index = item.get('order')
        db.session.commit()
        return jsonify({'status': 'success'})
    except Exception as e: return jsonify({'status': 'error', 'message': str(e)})

@app.route('/delete_row/<int:id>')
@login_required
def delete_row(id):
    row = DataRow.query.get(id)
    if row and row.created_by == current_user.id:
        tid = row.table_id
        db.session.delete(row)
        db.session.commit()
        return redirect(url_for('index', table_id=tid))
    return redirect(url_for('index'))

@app.route('/print_row/<int:id>')
@login_required
def print_row(id):
    row = DataRow.query.get_or_404(id)
    if row.created_by != current_user.id: return "Không có quyền!", 403
    columns = TableColumn.query.filter_by(user_id=current_user.id, table_id=row.table_id).order_by(TableColumn.order_index).all()
    vn_time = datetime.utcnow() + timedelta(hours=7)
    return render_template('print_ticket.html', row=row, columns=columns, today=vn_time.strftime('%d/%m/%Y'))

# --- USER ADMIN ROUTES ---
@app.route('/admin/toggle_role/<int:user_id>')
@login_required
@admin_required
def toggle_role(user_id):
    if user_id == current_user.id: return redirect(url_for('index'))
    user = User.query.get(user_id)
    if user:
        user.role = 'user' if user.role == 'admin' else 'admin'
        db.session.commit()
        flash('Đổi quyền thành công!', 'success')
    return redirect(url_for('index'))

@app.route('/admin/reset_password/<int:user_id>', methods=['POST'])
@login_required
@admin_required
def admin_reset_password(user_id):
    user = User.query.get(user_id)
    if user:
        user.password = generate_password_hash('123456')
        db.session.commit()
        flash('Reset mật khẩu thành công!', 'success')
    return redirect(url_for('index'))

@app.route('/admin/delete_user/<int:user_id>')
@login_required
@admin_required
def admin_delete_user(user_id):
    user = User.query.get(user_id)
    if user and user.role != 'admin':
        DataRow.query.filter_by(created_by=user.id).delete()
        TableColumn.query.filter_by(user_id=user.id).delete()
        db.session.delete(user)
        db.session.commit()
        flash('Đã xóa tài khoản!', 'success')
    return redirect(url_for('index'))

@app.route('/change_password', methods=['POST'])
@login_required
def change_password():
    new_pass = request.form.get('new_password')
    if new_pass:
        current_user.password = generate_password_hash(new_pass)
        db.session.commit()
        flash('Đổi mật khẩu thành công!', 'success')
    return redirect(url_for('index'))

@app.route('/login', methods=['GET', 'POST'])
def login():
    if request.method == 'POST':
        user = User.query.filter_by(username=request.form.get('username')).first()
        if request.form.get('action') == 'register':
            if user: flash('Trùng tên rồi!', 'error')
            else:
                role = 'admin' if request.form.get('username') == 'admin' else 'user'
                db.session.add(User(username=request.form.get('username'), password=generate_password_hash(request.form.get('password')), role=role))
                db.session.commit()
                flash('Tạo xong!', 'success')
        elif user and check_password_hash(user.password, request.form.get('password')):
            login_user(user)
            return redirect(url_for('index'))
        else: flash('Sai mật khẩu!', 'error')
    return render_template('login.html')

@app.route('/logout')
@login_required
def logout(): logout_user(); return redirect(url_for('login'))

@app.route('/ping_db')
def ping_db():
    try:
        count = User.query.count()
        return f"Hello Robot! {count}", 200
    except: return "Error", 500

with app.app_context(): db.create_all()

if __name__ == '__main__':
    app.run(host='0.0.0.0', port=5000, debug=True)